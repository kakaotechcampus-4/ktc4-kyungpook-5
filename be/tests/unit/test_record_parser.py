"""기록 원본 파일을 parse_result로 바꾸는 파서를 검증한다.

parse_result는 AI 색인 입력(IndexRecordRequest.blocks)과 같은 형식이다(#70, PR #88).
"""

import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.errors import DependencyError
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.ai import IndexRecordRequest
from app.core.enums import RecordFileType, RecordParseErrorReason
from app.services import record_parser
from app.services.record_parser import RecordParseError, _cell_text, parse_record


def _error(content: bytes, file_type: RecordFileType) -> RecordParseError:
    with pytest.raises(RecordParseError) as error:
        parse_record(content, file_type)
    return error.value


@pytest.mark.parametrize("file_type", [RecordFileType.HWP, RecordFileType.IMAGE])
def test_지원하지_않는_형식은_UNSUPPORTED_FORMAT(file_type):
    """PDF·CSV·XLSX만 받는다(#70). enum 값은 남아 있어도 파싱하지 않는다."""
    error = _error(b"anything", file_type)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert error.params == {}
    assert file_type.value in error.debug_message


def _blocks(content: bytes, file_type: RecordFileType) -> list[dict]:
    return parse_record(content, file_type)["blocks"]


def _texts(content: bytes, file_type: RecordFileType) -> list[str]:
    return [block["text"] for block in _blocks(content, file_type)]


def _labels(content: bytes, file_type: RecordFileType) -> list[str]:
    return [block["location"]["label"] for block in _blocks(content, file_type)]


CSV = RecordFileType.CSV


def test_CSV_행마다_열_이름과_값을_잇는다():
    """PR #88 예시 기록(2024 하반기 장부.csv)과 같은 모양이어야 한다."""
    content = "날짜,내용,입금,출금\n2024-10-02,가을 MT 회비,\"1,280,000\",\n".encode()

    assert _blocks(content, CSV) == [
        {
            "text": "날짜: 2024-10-02, 내용: 가을 MT 회비, 입금: 1,280,000, 출금:",
            "location": {"label": "2행"},
        }
    ]


def test_CSV_빈_줄은_건너뛰고_실제_줄_번호를_쓴다():
    content = "날짜,금액\n2025-03-03,45000\n\n2025-03-04,45000\n".encode()

    assert _labels(content, CSV) == ["2행", "4행"]


def test_CSV_CP949도_읽는다():
    """한국어 엑셀에서 저장한 CSV는 대개 CP949다."""
    content = "날짜,금액\n2025-03-03,45000\n".encode("cp949")

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 금액: 45000"]


def test_CSV_UTF8_BOM을_열_이름에_남기지_않는다():
    content = "\ufeff날짜,금액\n2025-03-03,45000\n".encode()

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 금액: 45000"]


def test_CSV_인코딩을_알_수_없으면_UNSUPPORTED_FORMAT():
    error = _error(b"\x80\x81\xff", CSV)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert "인코딩" in error.debug_message


def test_CSV_열_이름보다_긴_행은_열N으로_부른다():
    content = "날짜,금액\n2025-03-03,45000,메모\n".encode()

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 금액: 45000, 열3: 메모"]


def test_CSV_짧은_행은_빠진_열을_빈_값으로_적는다():
    content = "날짜,금액,비고\n2025-03-03,45000\n".encode()

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 금액: 45000, 비고:"]


def test_CSV_열_이름_칸이_비면_열N으로_부른다():
    content = "날짜,,비고\n2025-03-03,45000,\n".encode()

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 열2: 45000, 비고:"]


def test_CSV_앞뒤_공백을_지운다():
    content = " 날짜 , 금액 \n 2025-03-03 , 45000 \n".encode()

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 금액: 45000"]


@pytest.mark.parametrize(
    "content",
    [
        ('날짜,메모\n2025-03-03,"' + "x" * 140_000 + "\n").encode(),
        "날짜,금액\r2025-03-03,45000\r".encode(),
    ],
    ids=["필드_길이_초과", "CR만_쓰는_줄바꿈"],
)
def test_CSV_형식이_깨지면_UNSUPPORTED_FORMAT(content):
    """읽는 도중 csv 모듈이 거부하는 파일. 서버 오류가 아니라 사유로 알린다."""
    error = _error(content, CSV)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert "CSV" in error.debug_message


def test_CSV_셀의_NUL_문자를_지운다():
    """PostgreSQL JSONB는 \\u0000이 든 값을 저장하지 못한다(#42 parse_result 저장)."""
    content = "날짜,메모\n2025-03-03,봄\x00 MT\n,\x00\n".encode()

    assert _texts(content, CSV) == ["날짜: 2025-03-03, 메모: 봄 MT"]


def test_CSV_열_이름만_있으면_블록이_없다():
    """빈 기록은 #71 index_record가 '읽을 내용 없음'으로 처리한다."""
    assert _blocks("날짜,금액\n".encode(), CSV) == []


XLSX = RecordFileType.XLSX


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    """시트 이름과 행 목록으로 XLSX를 만든다. 빈 목록은 빈 행이다."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_XLSX_행마다_시트_이름과_행_번호를_붙인다():
    """PR #88 예시 기록(2025 상반기 장부.xlsx)과 같은 모양이어야 한다(숫자 쉼표 제외)."""
    content = _xlsx(
        {
            "회비 수납 시트": [
                ["날짜", "이름", "금액", "비고"],
                [datetime(2025, 3, 3), "참가자01", 45000, "봄 MT 회비"],
            ]
        }
    )

    assert _blocks(content, XLSX) == [
        {
            "text": "날짜: 2025-03-03, 이름: 참가자01, 금액: 45000, 비고: 봄 MT 회비",
            "location": {"label": "회비 수납 시트 · 2행"},
        }
    ]


def test_XLSX_시트_순서대로_읽는다():
    content = _xlsx(
        {
            "회비 수납": [["이름", "금액"], ["참가자01", 45000]],
            "지출": [["항목", "금액"], ["숙소", 550000]],
        }
    )

    assert _labels(content, XLSX) == ["회비 수납 · 2행", "지출 · 2행"]


def test_XLSX_위쪽_빈_행이_있어도_실제_행_번호를_쓴다():
    content = _xlsx({"장부": [[], [], ["날짜", "금액"], ["2025-03-03", 45000]]})

    assert _labels(content, XLSX) == ["장부 · 4행"]


def test_XLSX_중간_빈_행은_건너뛴다():
    content = _xlsx({"장부": [["이름"], ["참가자01"], [], ["참가자02"]]})

    assert _labels(content, XLSX) == ["장부 · 2행", "장부 · 4행"]


def test_XLSX_이름_없는_열은_값이_있을_때만_넣는다():
    """읽기 전용 모드는 행을 최대 열 수까지 None으로 채운다."""
    content = _xlsx(
        {
            "장부": [
                ["날짜", "금액"],
                ["2025-03-03", 45000, "추가 메모"],
                ["2025-03-04", 45000],
            ]
        }
    )

    assert _texts(content, XLSX) == [
        "날짜: 2025-03-03, 금액: 45000, 열3: 추가 메모",
        "날짜: 2025-03-04, 금액: 45000",
    ]


def _with_dimension(content: bytes, ref: str) -> bytes:
    """시트 XML의 크기 정보(<dimension ref>)를 바꾼다. 엑셀이 아닌 도구가 잘못 적는 경우를 흉내 낸다."""
    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"', f'<dimension ref="{ref}"'.encode(), data)
            target.writestr(item, data)
    return buffer.getvalue()


@pytest.mark.parametrize("ref", ["A1", "A1:B2"])
def test_XLSX_시트_크기_정보가_틀려도_모든_행과_열을_읽는다(ref):
    """읽기 전용 모드는 크기 정보만큼만 읽어, 틀리면 행·열이 오류 없이 빠진다."""
    content = _with_dimension(
        _xlsx(
            {
                "장부": [
                    ["날짜", "내용", "금액"],
                    ["2025-03-03", "회비", 45000],
                    ["2025-03-04", "회비", 45000],
                ]
            }
        ),
        ref,
    )

    assert _texts(content, XLSX) == [
        "날짜: 2025-03-03, 내용: 회비, 금액: 45000",
        "날짜: 2025-03-04, 내용: 회비, 금액: 45000",
    ]


def test_XLSX_데이터_없는_시트는_블록이_없다():
    content = _xlsx({"빈 시트": [], "열 이름만": [["날짜", "금액"]]})

    assert _blocks(content, XLSX) == []


@pytest.mark.parametrize(
    "content",
    [b"not a xlsx", None],
    ids=["zip_아님", "xlsx_아닌_zip"],
)
def test_XLSX_열_수_없으면_UNSUPPORTED_FORMAT(content):
    if content is None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("readme.txt", "x")
        content = buffer.getvalue()

    error = _error(content, XLSX)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert "열 수 없습니다" in error.debug_message


@pytest.mark.parametrize(
    ("value", "text"),
    [
        (None, ""),
        ("  봄 MT  ", "봄 MT"),
        (45000, "45000"),
        (45000.0, "45000"),
        (0.15, "0.15"),
        (49500.00000000001, "49500"),
        (datetime(2025, 3, 3), "2025-03-03"),
        (datetime(2025, 3, 3, 14, 30), "2025-03-03 14:30"),
        (date(2025, 3, 3), "2025-03-03"),
        (time(9, 5), "09:05"),
    ],
)
def test_표시_형식이_없는_셀은_원래_값으로_적는다(value, text):
    """일반 형식 셀. 연도·번호 열이 2,025가 되지 않고, 계산 오차만 정리한다."""
    assert _cell_text(value) == text


def _xlsx_formatted(cells: list[tuple[object, str]]) -> bytes:
    """(값, 셀 표시 형식) 목록을 열 이름 아래 한 행에 쓴 XLSX."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "장부"
    sheet.append([f"열{index}" for index in range(1, len(cells) + 1)])
    for column, (value, number_format) in enumerate(cells, start=1):
        cell = sheet.cell(row=2, column=column, value=value)
        cell.number_format = number_format
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("value", "number_format", "text"),
    [
        (45000, "#,##0", "45,000"),
        (2025, "General", "2025"),
        # 엑셀이 저장해 둔 계산값(0.1*3*165000)
        (49500.00000000001, "#,##0", "49,500"),
        (49500.00000000001, "General", "49500"),
        (1234.5, "#,##0.00", "1,234.50"),
        (45000, '"₩"#,##0', "45,000"),
    ],
)
def test_XLSX_셀_표시_형식의_쉼표와_소수_자릿수를_따른다(value, number_format, text):
    """엑셀 화면에 보이는 값과 같아야 같은 장부의 XLSX·CSV 결과가 같아진다."""
    [row] = _texts(_xlsx_formatted([(value, number_format)]), XLSX)

    assert row == f"열1: {text}"


def test_XLSX_시트_XML이_손상되면_UNSUPPORTED_FORMAT():
    """파일은 열리지만 행을 읽는 도중 깨진 경우."""
    source = zipfile.ZipFile(io.BytesIO(_xlsx({"장부": [["이름"], ["참가자01"]]})))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</sheetData>", b'<row r="9"><c r="A9"><v>1</v></c>')
            target.writestr(item, data)

    error = _error(buffer.getvalue(), XLSX)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert "XLSX" in error.debug_message


PDF = RecordFileType.PDF
SAMPLE_PDF = Path(__file__).parents[1] / "fixtures" / "records" / "sample_report.pdf"

_HELVETICA = DictionaryObject(
    {
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    }
)


def _text_pdf(*pages: str) -> bytes:
    """페이지마다 영문 한 줄을 쓴 PDF. 빈 문자열은 빈 페이지다."""
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(300, 300)
        if text:
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): _HELVETICA})}
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode("latin-1"))
            page.replace_contents(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _encrypted_pdf(user_password: str, algorithm: str) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_text_pdf("Restricted page"))))
    writer.encrypt(user_password, "owner-password", algorithm=algorithm)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_PDF_쪽마다_블록을_만든다():
    assert _blocks(_text_pdf("First page", "Second page"), PDF) == [
        {"text": "First page", "location": {"label": "1쪽"}},
        {"text": "Second page", "location": {"label": "2쪽"}},
    ]


def test_PDF_빈_페이지는_건너뛰고_쪽_번호는_유지한다():
    assert _labels(_text_pdf("First page", "", "Third page"), PDF) == ["1쪽", "3쪽"]


def test_PDF_모든_페이지가_비면_IMAGE_NOT_READABLE():
    """텍스트 층이 없는 스캔 문서. 화면은 '원본 문서로 다시 올리기'를 안내한다."""
    error = _error(_text_pdf("", ""), PDF)

    assert error.reason is RecordParseErrorReason.IMAGE_NOT_READABLE


def test_열람_암호가_있는_PDF는_UNSUPPORTED_FORMAT():
    error = _error(_encrypted_pdf("user-password", "RC4-128"), PDF)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert "암호" in error.debug_message


def test_권한_암호만_있는_PDF는_읽는다():
    """편집만 막은 PDF는 빈 비밀번호로 풀린다."""
    assert _texts(_encrypted_pdf("", "RC4-128"), PDF) == ["Restricted page"]


def test_AES_암호_해제에_라이브러리가_없으면_UNSUPPORTED_FORMAT(monkeypatch):
    """cryptography를 넣지 않아 AES PDF는 읽지 못한다. 서버 오류가 아니라 사유로 알린다."""

    def needs_cryptography(stream):
        raise DependencyError("cryptography>=3.1 is required for AES algorithm")

    monkeypatch.setattr(record_parser, "PdfReader", needs_cryptography)

    error = _error(b"%PDF-1.7", PDF)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT


@pytest.mark.parametrize("content", [b"not a pdf", b"", b"%PDF-1.4\n garbage"])
def test_PDF_열_수_없으면_UNSUPPORTED_FORMAT(content):
    error = _error(content, PDF)

    assert error.reason is RecordParseErrorReason.UNSUPPORTED_FORMAT
    assert "열 수 없습니다" in error.debug_message


def test_오류_메시지에_파일_내용을_넣지_않는다():
    error = _error(b"%PDF-1.4 secret-content-1234", PDF)

    assert "열 수 없습니다" in error.debug_message
    assert "secret-content-1234" not in error.debug_message


def _hangul_ratio(text: str) -> float:
    letters = [char for char in text if not char.isspace()]
    hangul = [char for char in letters if "\uac00" <= char <= "\ud7a3"]
    return len(hangul) / len(letters)


def test_샘플_PDF에서_한글을_추출한다():
    """한컴오피스 2022에서 내보낸 PDF(함초롬 글꼴). 글자 정보가 빠진 PDF는 텍스트가 깨진다."""
    blocks = _blocks(SAMPLE_PDF.read_bytes(), PDF)

    assert [block["location"]["label"] for block in blocks] == ["1쪽", "2쪽", "3쪽"]
    for block in blocks:
        assert "\ufffd" not in block["text"]
        assert _hangul_ratio(block["text"]) > 0.3


def test_샘플_PDF의_표는_숫자가_남는다():
    """2쪽 예산 표가 텍스트로 풀려도 금액 숫자는 남아야 한다(#72 평가 대상)."""
    page_two = _blocks(SAMPLE_PDF.read_bytes(), PDF)[1]["text"]

    assert any(char.isdigit() for char in page_two)


@pytest.mark.parametrize(
    ("content", "file_type"),
    [
        ("날짜,금액\n2025-03-03,45000\n".encode(), CSV),
        (_xlsx({"장부": [["날짜", "금액"], ["2025-03-03", 45000]]}), XLSX),
        (_text_pdf("First page"), PDF),
    ],
    ids=["CSV", "XLSX", "PDF"],
)
def test_파서_출력을_색인_요청으로_넘길_수_있다(content, file_type):
    """parse_result를 그대로 AI 색인 입력으로 넘긴다(PR #88 contracts.md)."""
    request = IndexRecordRequest(
        record_id="rec_test",
        file_name="기록",
        file_type=file_type,
        blocks=parse_record(content, file_type)["blocks"],
    )

    assert request.blocks
