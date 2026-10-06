"""업로드한 기록 원본(PDF·CSV·XLSX)을 parse_result로 바꾼다.

parse_result는 Record.parse_result에 저장하고 AI 색인 입력(IndexRecordRequest.blocks)으로
그대로 넘긴다(#70). 블록 순서가 블록 번호이며 빈 블록은 넣지 않는다.

    {"blocks": [{"text": "...", "location": {"label": "..."}}]}

위치 라벨은 원본의 물리적 위치만 적는다. PDF "N쪽" / XLSX "{시트 이름} · N행" / CSV "N행"이며
N은 원본의 실제 쪽·행 번호다. 제목·조항·행 범위 같은 논리적 위치는 청킹(#71)이 청크 라벨에 붙인다
(PR #90 리뷰 합의).
"""

import csv
import io
import re
import zipfile
from collections.abc import Callable, Iterable
from datetime import date, datetime, time
from typing import Any
from xml.etree.ElementTree import ParseError as XMLParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import DependencyError, PdfReadError

from app.core.enums import RecordFileType, RecordParseErrorReason

ParseResult = dict[str, list[dict[str, Any]]]


class RecordParseError(Exception):
    """파싱 실패. reason은 화면 문구를 고르는 코드, debug_message는 로그용 원문이다.

    parse_error로 저장하는 형식(키 표기)은 호출하는 쪽(#42)이 정한다.
    """

    def __init__(
        self,
        reason: RecordParseErrorReason,
        debug_message: str,
        params: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(debug_message)
        self.reason = reason
        self.debug_message = debug_message
        self.params = params or {}


def parse_record(content: bytes, file_type: RecordFileType) -> ParseResult:
    """원본 파일을 parse_result로 바꾼다.

    예상하지 못한 예외는 원인을 단정하지 않도록 감싸지 않는다(AI-RET-04).
    """
    parser = _PARSERS.get(file_type)
    if parser is None:
        raise RecordParseError(
            RecordParseErrorReason.UNSUPPORTED_FORMAT,
            f"지원하지 않는 형식입니다: {file_type.value}",
        )
    return {"blocks": parser(content)}


def _block(text: str, label: str) -> dict[str, Any]:
    return {"text": text, "location": {"label": label}}


def _parse_pdf(content: bytes) -> list[dict[str, Any]]:
    try:
        reader = PdfReader(io.BytesIO(content))
        # 편집만 막은 PDF는 빈 비밀번호로 풀린다. 열람 암호가 있으면 풀리지 않는다.
        if reader.is_encrypted and not reader.decrypt(""):
            # [합의 필요] 화면 문구가 "지원하지 않는 형식"이라 업로드 시 거절(#42)할지
            # 별도 사유를 둘지 FE와 협의 중이다(PR #90).
            raise RecordParseError(
                RecordParseErrorReason.UNSUPPORTED_FORMAT, "열람 암호가 걸린 PDF입니다"
            )
        pages = [_clean(page.extract_text() or "") for page in reader.pages]
    except (PdfReadError, DependencyError) as error:
        # DependencyError: AES 암호 해제에 cryptography가 필요하다. 의존성을 늘리지 않기로 해 지원하지 않는다.
        raise RecordParseError(
            RecordParseErrorReason.UNSUPPORTED_FORMAT,
            f"PDF 파일을 열 수 없습니다: {type(error).__name__}",
        ) from error

    blocks = [
        _block(text, f"{number}쪽") for number, text in enumerate(pages, start=1) if text
    ]
    if not blocks:
        raise RecordParseError(
            RecordParseErrorReason.IMAGE_NOT_READABLE,
            "모든 페이지에 텍스트가 없습니다 (스캔 문서로 보임)",
        )
    return blocks


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    try:
        # data_only: 수식 대신 엑셀이 저장해 둔 계산값을 읽는다.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, InvalidFileException) as error:
        raise RecordParseError(
            RecordParseErrorReason.UNSUPPORTED_FORMAT,
            f"XLSX 파일을 열 수 없습니다: {type(error).__name__}",
        ) from error
    try:
        blocks: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            # 읽기 전용 모드는 파일에 적힌 시트 크기만큼만 읽는다. 엑셀이 아닌 도구가 크기를
            # 잘못 적으면 행·열이 오류 없이 빠지므로 저장된 크기를 버리고 끝까지 읽는다.
            sheet.reset_dimensions()
            # 읽기 전용 모드의 iter_rows는 min_row와 상관없이 1행부터 준다. 1부터 세야 실제 행 번호다.
            # 표시 형식(number_format)을 보려고 값이 아니라 셀로 읽는다.
            rows = (
                (number, [_cell_text(cell.value, cell.number_format) for cell in row])
                for number, row in enumerate(sheet.iter_rows(min_row=1), start=1)
            )
            blocks.extend(
                _table_blocks(rows, lambda number, title=sheet.title: f"{title} · {number}행")
            )
        return blocks
    except (XMLParseError, zipfile.BadZipFile) as error:
        # 파일은 열렸지만 시트를 읽는 도중 깨진 경우(시트 XML 손상, 압축 데이터 손상).
        raise RecordParseError(
            RecordParseErrorReason.UNSUPPORTED_FORMAT,
            f"XLSX 시트를 읽을 수 없습니다: {type(error).__name__}",
        ) from error
    finally:
        workbook.close()


# 한국어 엑셀에서 저장한 CSV는 대개 CP949다. BOM이 붙은 UTF-8도 흔하다.
_CSV_ENCODINGS = ("utf-8-sig", "cp949")


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    for encoding in _CSV_ENCODINGS:
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise RecordParseError(
            RecordParseErrorReason.UNSUPPORTED_FORMAT,
            "CSV 인코딩을 읽을 수 없습니다 (UTF-8·CP949 아님)",
        )
    # 레코드 순번을 행 번호로 쓴다. 따옴표 안 줄바꿈이 있으면 실제 줄 번호와 다를 수 있다.
    rows = (
        (number, [_cell_text(cell) for cell in row])
        for number, row in enumerate(csv.reader(io.StringIO(text)), start=1)
    )
    try:
        return _table_blocks(rows, lambda number: f"{number}행")
    except csv.Error as error:
        # CR만 쓰는 줄바꿈, 필드 길이 초과처럼 읽는 도중 csv 모듈이 거부한 경우.
        raise RecordParseError(
            RecordParseErrorReason.UNSUPPORTED_FORMAT,
            f"CSV 형식을 읽을 수 없습니다: {type(error).__name__}",
        ) from error


def _table_blocks(
    rows: Iterable[tuple[int, list[str]]], label: Callable[[int], str]
) -> list[dict[str, Any]]:
    """처음으로 비어 있지 않은 행을 열 이름으로, 이후 행을 하나씩 블록으로 만든다.

    제목 줄이 위에 있는 장부는 고려하지 않은 규칙이다. 장부 파싱 규칙이 정해지면 바꾼다.
    """
    header: list[str] | None = None
    blocks: list[dict[str, Any]] = []
    for number, values in rows:
        if not any(values):
            continue
        if header is None:
            header = values
            continue
        blocks.append(_block(_row_text(header, values), label(number)))
    return blocks


def _row_text(header: list[str], values: list[str]) -> str:
    """'열 이름: 값'을 잇는다. 행만 떼어 봐도 무슨 값인지 알 수 있게 한다(PR #88).

    이름 없는 열은 '열N'으로 부르고, 이름도 값도 없는 칸은 뺀다.
    """
    pairs: list[str] = []
    for index in range(max(len(header), len(values))):
        name = header[index] if index < len(header) else ""
        value = values[index] if index < len(values) else ""
        if not name and not value:
            continue
        name = name or f"열{index + 1}"
        pairs.append(f"{name}: {value}" if value else f"{name}:")
    return ", ".join(pairs)


# "#,##0", "#,##0.00", '"₩"#,##0'처럼 천 단위 쉼표가 있는 표시 형식. 소수 자릿수는 뒤의 0 개수다.
_THOUSANDS_FORMAT = re.compile(r"#,##0(?:\.(0+))?")


def _cell_text(value: object, number_format: str | None = None) -> str:
    """셀 값을 엑셀 화면에 보이는 대로 적는다(PR #90 리뷰 합의).

    천 단위 쉼표 형식만 따르고, 일반 형식 숫자는 원래 값으로 둔다. 연도·번호 열이
    2,025가 되지 않게 하기 위해서다. 같은 장부의 XLSX와 엑셀이 저장한 CSV가 같은 값이 된다.
    """
    if value is None:
        return ""
    # datetime은 date의 하위 타입이라 먼저 본다.
    if isinstance(value, datetime):
        if value.time() == time():
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    # bool은 int의 하위 타입이라 숫자 표기에서 뺀다.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        thousands = _THOUSANDS_FORMAT.search(number_format or "")
        if thousands:
            return f"{value:,.{len(thousands.group(1) or '')}f}"
        if isinstance(value, float):
            # 일반 형식의 계산값은 부동소수점 오차만 정리한다(49500.00000000001 → 49500).
            return f"{value:.15g}"
        return str(value)
    return _clean(str(value))


def _clean(text: str) -> str:
    """앞뒤 공백과 NUL 문자를 지운다. PostgreSQL JSONB는 \\u0000이 든 값을 저장하지 못한다(#42)."""
    return text.replace("\x00", "").strip()


_PARSERS: dict[RecordFileType, Callable[[bytes], list[dict[str, Any]]]] = {
    RecordFileType.PDF: _parse_pdf,
    RecordFileType.CSV: _parse_csv,
    RecordFileType.XLSX: _parse_xlsx,
}
